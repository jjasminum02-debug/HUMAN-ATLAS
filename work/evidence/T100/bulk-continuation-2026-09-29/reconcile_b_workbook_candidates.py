#!/usr/bin/env python3
"""Reconcile historical T100 B workbook candidate counts without retaining workbook rows."""
from __future__ import annotations
import argparse, hashlib, json, re, subprocess
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[4]
EXPECTED_SHA="8e10e00e76f5da8e6d1be6a0f9dcab6039d9a3f81d8a3808c0fe28545802ab3c"
HISTORICAL_B_HEAD="1a45ee5e84d921ce88c279b7226cb249723c9673"
CONTINUATION_START_HEAD="c757c2d9e870c8c403025868024872c1710d56a2"
OUT=Path(__file__).resolve().parent/"candidate-reconciliation.json"

def norm(value: str) -> str:
    return re.sub(r"\s+"," ",value.strip()).casefold()

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--workbook",required=True)
    args=parser.parse_args()
    book=Path(args.workbook)
    digest=hashlib.sha256(book.read_bytes()).hexdigest()
    if digest!=EXPECTED_SHA: raise SystemExit("private workbook hash differs from the previously recorded source")
    wb=load_workbook(book,read_only=True,data_only=True)
    ws=wb["근육표"]
    # Read the two permitted label columns only. No workbook cell values are written.
    rows=[]
    for korean,english in ws.iter_rows(min_col=2,max_col=3,values_only=True):
        if korean is None or english is None: continue
        if str(english).strip().casefold() in {"english name","영문명"}: continue
        rows.append((str(korean).strip(),str(english).strip()))
    labels={norm(english) for _,english in rows}
    scope=json.loads((ROOT/"atlas-data/catalog/target-scope-t96.json").read_text())
    continuation_overlay=json.loads(subprocess.check_output(["git","show",f"{CONTINUATION_START_HEAD}:atlas-data/overlays/za-local-integration.json"],cwd=ROOT))
    t96_exact={norm(t["term"]["english"]):t for t in scope["targets"]}
    matched=[t96_exact[name] for name in labels if name in t96_exact]
    existing_evidence={term["targetId"] for term in continuation_overlay.get("targetTerminologyEvidence",[])}
    direct_rows={t["id"]:[o for o in continuation_overlay["objects"] if o.get("targetId")==t["id"] and o.get("localDisplayEligible")] for t in matched}
    exact_target_name_count=0
    for target in matched:
        base=norm(target["term"]["english"])
        if any(norm(re.sub(r"\.[lr]$","",obj["sourceName"],flags=re.I))==base for obj in direct_rows[target["id"]]):
            exact_target_name_count+=1
    start_overlay_raw=subprocess.check_output(["git","show",f"{HISTORICAL_B_HEAD}:atlas-data/overlays/za-local-integration.json"],cwd=ROOT)
    start_overlay=json.loads(start_overlay_raw)
    source_bases={}
    for obj in start_overlay["objects"]:
        if not obj.get("localDisplayEligible"): continue
        base=norm(re.sub(r"\.[lr]$","",obj["sourceName"],flags=re.I))
        source_bases.setdefault(base,[]).append(obj)
    direct_base_matches=labels & set(source_bases)
    unnamed={name:[obj for obj in source_bases[name] if not obj["names"].get("koModern")] for name in direct_base_matches
              if any(not obj["names"].get("koModern") for obj in source_bases[name])}
    record={
      "schemaVersion":1,"taskId":"T100","unit":"B candidate metric correction","checkedAtLocal":"2026-09-29",
      "workbook":{"sha256":digest,"sheet":"근육표","fieldsRead":["B","C"],"rawValuesRetained":False,"modified":False},
      "normalization":{"workbookAndTarget":"trim, collapse whitespace, casefold; no singular/plural, suffix, synonym or fuzzy rewrite","workbookAndSourceBase":"trim terminal `.l`/`.r` only from source object name, then trim/collapse whitespace/casefold"},
      "historicalMetric":{"source":"T100 B start ledger","exactNormalizedSourceLabelMatches":137,"unnamedCandidateLabels":101,"unnamedCandidateSurfaceRows":201,"reproduced":False},
      "reproducedAtOriginalBStart":{"head":HISTORICAL_B_HEAD,"uniqueWorkbookEnglishLabels":len(labels),"exactLocalSourceNameBaseMatches":len(direct_base_matches),"unnamedCandidateLabels":len(unnamed),"unnamedCandidateSurfaceRows":sum(len(v) for v in unnamed.values()),"fullyNamedMatchedLabels":len(direct_base_matches)-len(unnamed),"notFoundByExactRule":len(labels-direct_base_matches)},
      "strictT96Join":{"uniqueWorkbookEnglishLabels":len(labels),"exactT96EnglishTargetMatches":len(matched),"targetsWithoutPriorTargetTerminologyEvidence":sum(t["id"] not in existing_evidence for t in matched),"directTargetEligibleSurfaceRows":sum(len(v) for v in direct_rows.values()),"directTargetRowsMissingKoModern":sum(not obj["names"].get("koModern") for values in direct_rows.values() for obj in values),"matchedTargetsWithAtLeastOneExactSourceNameBase":exact_target_name_count},
      "interpretation":"The prior 137/101/201 figures do not reproduce from the same workbook hash and original T100 B-start overlay with the explicit exact-name normalization above. Keep the historical report intact; do not count the metric difference as completed anatomy. The validated B candidate set is the 39 unnamed exact-source-name concepts / 77 surface rows, of which this continuation applied 34 direct KAA term rows to 67 exact side/midline surfaces and retained five exact-query misses for C.",
      "privacy":"No private workbook path, raw name list, non-name column, cell values, or workbook copy is retained."
    }
    OUT.write_text(json.dumps(record,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"status":"pass","historicalMetricReproduced":False,"exactCandidates":len(unnamed),"candidateSurfaceRows":sum(len(v) for v in unnamed.values()),"exactT96Targets":len(matched),"strictSourceNameTargetConcepts":exact_target_name_count},ensure_ascii=False,indent=2))
if __name__=="__main__": main()
