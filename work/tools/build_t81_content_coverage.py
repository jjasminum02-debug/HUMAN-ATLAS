#!/usr/bin/env python3
"""Import the private T81 workbook as candidate-only evidence and build coverage.

Raw workbook values and extracted candidate payloads are written only below the
ignored private-sources/ directory. Tracked outputs contain row hashes,
dispositions, and a learner-safe source-scoped pointer to already verified
attachment claims. This tool does not create canonical bindings or claims.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
PRIVATE_SOURCE = ROOT / "private-sources/t81/muscle-attachment-notes.xlsx"
PRIVATE_EXTRACT = ROOT / "private-sources/t81/candidate-extract.json"
EVIDENCE = ROOT / "work/evidence/T81"
CONTENT_PATH = ROOT / "atlas-data/terminology/learner-structure-source-content.json"
WORKBOOK_AUDIT = ROOT / "work/evidence/2026-09-28-workbook-observation-plan/workbook-audit.json"

CANONICAL_TARGETS = {
    "TA2:2660": "HA-M-000002",  # soleus
    "TA2:2644": "HA-M-000003",  # tibialis anterior
    "TA2:2666": "HA-M-000004",  # tibialis posterior
    "TA2:2652": "HA-M-000005",  # fibularis longus
    "TA2:2653": "HA-M-000006",  # fibularis brevis
    "TA2:2658": "HA-P-000001",  # lateral head of gastrocnemius
    "TA2:2659": "HA-P-000002",  # medial head of gastrocnemius
}

# These workbook labels describe broad lecture groupings, so they only filter
# exact English-name candidates. They never establish a claim or canonical ID.
WORKBOOK_REGION_FAMILIES = {
    "머리·얼굴": {"head"},
    "씹기·목·인두": {"head", "neck"},
    "목·척주": {"neck", "back"},
    "가슴우리·배·골반": {"thorax", "abdomen-lumbar", "pelvis-perineum"},
    "팔이음뼈·어깨·위팔": {"shoulder-scapular", "upper-limb"},
    "아래팔·손": {"upper-limb"},
    "볼기·넙다리": {"gluteal-hip", "thigh"},
    "종아리·발": {"leg", "foot"},
}


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def normalize_exact(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).casefold().replace("\u00a0", " ")
    return " ".join(text.split())


def json_read(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def candidate_fields(workbook_path: Path) -> tuple[list[dict[str, Any]], dict[str, int]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    if "근육표" not in workbook.sheetnames or "표기 안내" not in workbook.sheetnames:
        raise ValueError("Expected the original 근육표 and 표기 안내 sheets.")
    sheet = workbook["근육표"]
    headers = [sheet.cell(5, column).value for column in range(1, 9)]
    expected = ["부위", "근육명", "English name", "기시", "정지", "운동신경", "감각신경·고유감각 경로", "작용(한자어 용어)"]
    if headers != expected:
        raise ValueError("Workbook header row changed; do not reinterpret the private workbook.")
    rows: list[dict[str, Any]] = []
    for row_number in range(6, 263):
        values = [sheet.cell(row_number, column).value for column in range(1, 9)]
        if not any(value is not None for value in values):
            continue
        row_bytes = json.dumps(values, ensure_ascii=False, separators=(",", ":"), default=str).encode()
        rows.append({
            "row": row_number,
            "region": values[0],
            "koreanName": values[1],
            "englishName": values[2],
            "originCandidate": values[3],
            "insertionCandidate": values[4],
            "motorNerveCandidate": values[5],
            "sensoryProprioceptionCandidate": values[6],
            "actionHanCharacterCandidate": values[7],
            "rowSha256": sha256_bytes(row_bytes),
            "rowValues": values,
        })
    link_count = 0
    comment_count = 0
    formula_count = 0
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row or 1, max_col=sheet.max_column or 8):
        for cell in row:
            link_count += int(cell.hyperlink is not None)
            comment_count += int(cell.comment is not None)
            formula_count += int(isinstance(cell.value, str) and cell.value.startswith("="))
    audit = {
        "sheetRows": len(rows),
        "sourceHyperlinks": link_count,
        "cellComments": comment_count,
        "formulas": formula_count,
        "hasExternalLinks": bool(getattr(workbook, "_external_links", [])),
    }
    return rows, audit


def target_rows() -> dict[str, dict[str, Any]]:
    scope = json_read(ROOT / "atlas-data/catalog/target-scope-t96.json")
    return {row["id"]: row for row in scope["targets"]}


def build(workbook_path: Path) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Private workbook missing: {workbook_path}")
    raw = workbook_path.read_bytes()
    workbook_hash = sha256_bytes(raw)
    expected_audit = json_read(WORKBOOK_AUDIT)
    expected = expected_audit["sha256"]
    if workbook_hash != expected:
        raise ValueError(f"Workbook SHA256 differs from the preserved audit: {workbook_hash}")

    product_scope = json_read(ROOT / "work/product-scope.json")
    compiled = json_read(ROOT / "atlas-data/source-cache/datasets/za/compiled/manifest.json")
    target_scope = target_rows()
    names = json_read(ROOT / "atlas-data/terminology/learning-names.json")
    ai_overlay = json_read(ROOT / "atlas-data/terminology/ai-evidence-overlay.json")
    if product_scope["denominators"]["targets"] != 542 or product_scope["denominators"]["memberships"] != 563 or product_scope["denominators"]["regions"] != 12:
        raise ValueError("T96 product denominators changed; do not rebase this T81 import.")
    if product_scope["denominators"]["existingCanonicalHaBindings"] != 130:
        raise ValueError("Existing HA binding denominator changed; stop and reconcile first.")

    muscle_instances = {row["sourceKey"]: row for row in compiled["instances"] if row["kind"] == "muscle_surface_or_part"}
    supported = [row for row in product_scope["supportedStructures"] if row["sourceKey"] in muscle_instances]
    by_concept: dict[tuple[str, tuple[str, ...]], list[dict[str, Any]]] = defaultdict(list)
    for row in supported:
        source = muscle_instances[row["sourceKey"]]
        data_name = source.get("dataName")
        if not data_name:
            raise ValueError(f"Supported muscle instance missing source dataName: {row['sourceKey']}")
        by_concept[(data_name, tuple(sorted(row["regionIds"])))].append(row)
    if len(supported) != 462 or len(by_concept) != 232:
        raise ValueError(f"Supported muscle scope drifted: surfaces={len(supported)}, concepts={len(by_concept)}")

    rows, workbook_audit = candidate_fields(workbook_path)
    if len(rows) != 257:
        raise ValueError(f"Workbook row scope drifted: {len(rows)}")
    workbook_by_english: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        workbook_by_english[normalize_exact(row["englishName"])].append(row)

    concept_by_english: dict[str, list[tuple[str, tuple[str, ...]]]] = defaultdict(list)
    for key in by_concept:
        concept_by_english[normalize_exact(key[0])].append(key)

    concept_candidates: dict[tuple[str, tuple[str, ...]], dict[str, Any]] = {}
    for key, concept_rows in by_concept.items():
        sources = [muscle_instances[row["sourceKey"]] for row in concept_rows]
        concept_candidates[key] = {
            "sourceDataName": key[0],
            "regionIds": list(key[1]),
            "sourceKeys": sorted(row["sourceKey"] for row in concept_rows),
            "sides": sorted({row.get("side") for row in concept_rows}, key=lambda value: str(value)),
            "targetIds": sorted({scope["targetId"] for row in concept_rows for scope in row.get("representedScopes", []) if isinstance(scope.get("targetId"), str)}),
            "scopeTypes": sorted({scope["scopeType"] for row in concept_rows for scope in row.get("representedScopes", []) if scope.get("scopeType")}),
            "sourceOnly": all(source["sourceOnly"] for source in sources),
        }

    row_dispositions: list[dict[str, Any]] = []
    exact_candidate_rows_by_concept: dict[tuple[str, tuple[str, ...]], list[int]] = defaultdict(list)
    disposition_counts: Counter[str] = Counter()
    for row in rows:
        english_key = normalize_exact(row["englishName"])
        candidates = concept_by_english.get(english_key, []) if english_key else []
        allowed_regions = WORKBOOK_REGION_FAMILIES.get(str(row["region"]), set())
        regional = [key for key in candidates if allowed_regions.intersection(key[1])]
        if len(regional) == 1:
            status = "exact_english_and_region_candidate_only_unverified"
            exact_candidate_rows_by_concept[regional[0]].append(row["row"])
        elif len(regional) > 1:
            status = "ambiguous_multiple_supported_source_concepts"
        elif candidates:
            status = "english_match_region_scope_conflict"
        else:
            status = "no_exact_match_in_current_supported_muscles"
        disposition_counts[status] += 1
        row_dispositions.append({
            "sheet": "근육표",
            "row": row["row"],
            "rowSha256": row["rowSha256"],
            "disposition": status,
            "candidateSupportedConceptKeys": [
                "|".join((key[0], *key[1])) for key in regional
            ],
            "candidateSupportedConcepts": [concept_candidates[key] for key in regional],
            "exactEnglishCandidatesOutsideWorkbookRegion": [
                concept_candidates[key] for key in candidates if key not in regional
            ],
            "candidateIsEvidence": False,
            "directRowSourceCitation": False,
            "workbookColumnsAreCandidateOnly": ["기시", "정지", "운동신경", "감각신경·고유감각 경로", "작용(한자어 용어)"],
            "actionColumnUse": "candidate_only_not_imported_into_learner_content",
        })

    name_entries = {entry["id"]: entry for entry in names["entries"]}
    ai_fields = {
        (entry["subjectId"], entry["field"]): entry
        for entry in ai_overlay["items"]
        if entry.get("field") in {"origin", "insertion"}
    }
    content_concepts: list[dict[str, Any]] = []
    production_records: list[dict[str, Any]] = []
    source_scoped_pointer_count = 0
    claim_coverage_counts: Counter[str] = Counter()
    for key, source_rows in sorted(by_concept.items()):
        data_name, region_ids = key
        source_records = [muscle_instances[row["sourceKey"]] for row in source_rows]
        source_keys = sorted(row["sourceKey"] for row in source_rows)
        target_ids = sorted({scope["targetId"] for row in source_rows for scope in row.get("representedScopes", []) if isinstance(scope.get("targetId"), str)})
        workbook_row_ids = sorted(exact_candidate_rows_by_concept.get(key, []))
        sides = sorted({row.get("side") for row in source_rows}, key=lambda value: str(value))

        mapped_subjects: list[tuple[str, str, str, dict[str, Any]]] = []
        for target_id, subject_id in CANONICAL_TARGETS.items():
            related = [(row, scope) for row in source_rows for scope in row.get("representedScopes", []) if scope["targetId"] == target_id and scope["relationKind"] == "normalized_exact_target_term"]
            if not related:
                continue
            target = target_scope.get(target_id)
            canonical = name_entries.get(subject_id)
            if not target or not canonical:
                raise ValueError(f"Missing target/canonical content for bounded projection {target_id}/{subject_id}")
            def strip_terminal_muscle(value: str) -> str:
                return re.sub(r"\s+muscle$", "", normalize_exact(value))
            if strip_terminal_muscle(data_name) != strip_terminal_muscle(target["term"]["english"]):
                continue
            if strip_terminal_muscle(canonical["english"]) != strip_terminal_muscle(target["term"]["english"]):
                continue
            scope_types = {scope["scopeType"] for _, scope in related}
            expected_kind = "explicit_part" if target["semanticKind"] == "muscle_part" else "whole_structure"
            if scope_types != {expected_kind}:
                raise ValueError(f"Wrong scope for {data_name} -> {target_id}: {scope_types}")
            if len(source_keys) != 2 or set(sides) != {"left", "right"}:
                raise ValueError(f"Expected explicit bilateral source surfaces for {target_id}: {source_keys}")
            mapped_subjects.append((target_id, subject_id, expected_kind, canonical))

        if len(mapped_subjects) > 1:
            raise ValueError(f"A supported source concept maps to multiple content subjects: {data_name}")
        subject_id = mapped_subjects[0][1] if mapped_subjects else None
        target_id = mapped_subjects[0][0] if mapped_subjects else None
        scope_type = mapped_subjects[0][2] if mapped_subjects else None
        field_states: dict[str, dict[str, Any]] = {}
        for field in ("origin", "insertion"):
            ai = ai_fields.get((subject_id, field)) if subject_id else None
            if not ai:
                field_states[field] = {"status": "no_verified_field_claim", "evidenceState": None, "subjectId": None}
            else:
                state = ai["evidenceState"]
                field_states[field] = {"status": "claim_available" if state != "conflicted" else "conflicted_not_asserted", "evidenceState": state, "subjectId": subject_id,
                                       "claimIds": [claim["id"] for claim in ai["claims"]]}
            claim_coverage_counts[field_states[field]["status"]] += 1
        for field in ("motorNerve", "sensoryProprioception"):
            field_states[field] = {"status": "no_verified_field_claim", "evidenceState": None, "subjectId": None}
            claim_coverage_counts[field_states[field]["status"]] += 1

        common = {
            "sourceDataName": data_name,
            "regionIds": list(region_ids),
            "sourceKeys": source_keys,
            "sourceSides": sides,
            "targetIds": target_ids,
            "exactWorkbookCandidateRows": workbook_row_ids,
            "workbookCandidatesAreClaims": False,
            "fieldDisposition": field_states,
        }
        content_concepts.append({
            **common,
            "conceptKey": f"za-source-name:{sha256_bytes(data_name.encode('utf-8'))[:20]}",
            "status": "source_identity_only",
            "canonicalContentSubjectId": subject_id,
            "targetContentScope": target_id,
            "scopeType": scope_type,
            "sourceOnly": all(source["sourceOnly"] for source in source_records),
            "canonicalLearnerBindingCreated": False,
            "humanReview": "not_performed",
            "publicRedistribution": "held",
        })
        if subject_id:
            source_scoped_pointer_count += len(source_keys)
        for source_row in source_rows:
            source = muscle_instances[source_row["sourceKey"]]
            if source.get("canonicalConceptId") is not None or source.get("learnerBinding") != "source_only_unbound":
                raise ValueError(f"T81 cannot change source-only canonical binding: {source['sourceKey']}")
            if source.get("humanReview") != "not_performed" or source.get("publicRedistribution") != "held":
                raise ValueError(f"T81 source preservation state changed: {source['sourceKey']}")
        production_records.append({
            "sourceKeys": source_keys,
            "targetId": target_id,
            "subjectId": subject_id,
            "scopeType": scope_type,
            "regionIds": list(region_ids),
            "sideSet": sides,
            "fieldDisposition": {key: {"status": value["status"], "evidenceState": value["evidenceState"]} for key, value in field_states.items()},
            "sourceOnly": True,
            "humanReview": "not_performed",
            "publicRedistribution": "held",
            "canonicalBindingCreated": False,
        })

    if len(production_records) != 232 or source_scoped_pointer_count != 14:
        raise ValueError(f"Unexpected learner content coverage: concepts={len(production_records)}, claim pointers={source_scoped_pointer_count}")

    raw_candidate_payload = {
        "schemaVersion": "t81-private-workbook-candidate-extract-v1",
        "sourceSha256": workbook_hash,
        "sourceSheet": "근육표",
        "sourceRows": [{key: value for key, value in row.items() if key != "rowSha256"} for row in rows],
        "notice": "Private candidate text only. Not claim evidence; not copied to app data or learner bundle.",
    }
    if PRIVATE_EXTRACT.exists():
        existing = json_read(PRIVATE_EXTRACT)
        if existing.get("sourceSha256") != workbook_hash:
            raise ValueError("Private extract exists for a different workbook; preserve it and stop.")
    else:
        write_json(PRIVATE_EXTRACT, raw_candidate_payload)

    content_data = {
        "schemaVersion": "learner-structure-source-content-v1",
        "revision": "t81-2026-10-01",
        "contract": {
            "scope": "source-scoped display pointer to existing field evidence only",
            "notCanonicalBinding": True,
            "newClaims": 0,
            "newGeometry": 0,
            "humanReview": "not_performed",
            "publicRedistribution": "held",
            "sourceOnly": True,
        },
        "records": sorted(production_records, key=lambda row: row["sourceKeys"][0]),
    }

    public_row_dispositions = {
        "schemaVersion": "t81-private-workbook-row-dispositions-v1",
        "source": {
            "originalPath": expected_audit["localOriginalPath"],
            "sha256": workbook_hash,
            "bytes": len(raw),
            "privateStagedPath": "private-sources/t81/muscle-attachment-notes.xlsx",
            "payloadIncluded": False,
            "rowSourceReferences": workbook_audit,
            "interpretation": "Workbook cells are candidate notes; missing row-level references and missing lecture PDFs prevent claim adoption without independent source verification.",
        },
        "sheet": "근육표",
        "rowsExpected": 257,
        "rowsProcessed": len(row_dispositions),
        "counts": dict(sorted(disposition_counts.items())),
        "rows": row_dispositions,
    }

    source_hashes = {
        path: sha256_file(ROOT / path)
        for path in [
            "work/product-scope.json",
            "atlas-data/source-cache/datasets/za/compiled/manifest.json",
            "atlas-data/catalog/target-scope-t96.json",
            "atlas-data/terminology/learning-names.json",
            "atlas-data/terminology/ai-evidence-overlay.json",
            "atlas-data/terminology/learning-structure-summaries.json",
            "work/evidence/T58/app-finish-2026-10-01/final-validation.json",
        ]
    }
    supported_coverage = {
        "schemaVersion": "t81-supported-muscle-content-coverage-v1",
        "asOf": date.today().isoformat(),
        "inputs": source_hashes,
        "denominators": {
            "targets": 542, "memberships": 563, "regions": 12,
            "existingCanonicalHaBindings": 130,
            "historical163": {"total": 163, "categories": {"6": 6, "20": 20, "135": 135, "2": 2}},
            "supportedMuscleSurfaces": len(supported),
            "uniqueSupportedMuscleSourceConcepts": len(content_concepts),
            "workbookRows": len(rows),
        },
        "supportedMuscles": content_concepts,
        "fieldDispositionCounts": dict(sorted(claim_coverage_counts.items())),
        "policy": {
            "workbookIsCandidateOnly": True,
            "workbookActionColumnNotImported": True,
            "motorAndSensoryProprioceptionSeparated": True,
            "noCanonicalBindingCreated": True,
            "noGeometryOrCoordinatesCreated": True,
            "humanReview": "not_performed",
            "publicRedistribution": "held",
        },
    }

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    write_json(EVIDENCE / "workbook-row-dispositions.json", public_row_dispositions)
    write_json(EVIDENCE / "supported-muscle-content-coverage.json", supported_coverage)
    write_json(CONTENT_PATH, content_data)
    return {
        "workbookSha256": workbook_hash,
        "workbookRows": len(rows),
        "workbookDispositionCounts": dict(disposition_counts),
        "muscleSurfaces": len(supported),
        "uniqueMuscleSourceConcepts": len(content_concepts),
        "supportedConceptCoverageRecords": len(production_records),
        "sourceScopedClaimPointers": source_scoped_pointer_count,
        "contentFieldDispositionCounts": dict(claim_coverage_counts),
        "workbookSourceAudit": workbook_audit,
        "sourceHashes": source_hashes,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=PRIVATE_SOURCE)
    args = parser.parse_args()
    result = build(args.source)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as error:
        print(f"T81 content coverage build failed: {error}", file=sys.stderr)
        sys.exit(1)
